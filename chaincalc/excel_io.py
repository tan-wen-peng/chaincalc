"""Excel 输入模板的生成与解析（openpyxl），以及 JSON 输入。"""

from __future__ import annotations

import json
import os
from typing import Any, Dict, List, Optional

from .model import Chain, Ring

PARAM_SHEET = "参数"
CHAIN_SHEET = "尺寸链"
GDT_SHEET = "GD&T"
HELP_SHEET = "说明"

CHAIN_HEADERS = [
    ("序号", 6), ("环名称", 22), ("类型", 10), ("基本尺寸mm", 14),
    ("上偏差mm", 12), ("下偏差mm", 12), ("分布系数k", 12),
    ("不对称系数e", 14), ("固定", 8), ("来源", 12), ("备注", 40),
]

GDT_HEADERS = [
    ("序号", 6), ("环名称", 22), ("几何特征", 14), ("规格值mm", 12),
    ("类型", 10), ("是否Ⓜ", 10), ("尺寸公差mm", 14), ("要素尺寸mm", 14),
    ("补偿方式", 12), ("作用", 12), ("备注", 40),
]

PARAM_ROWS = [
    ("项目名称", "示例：隔板-垫片-轴套轴向间隙链"),
    ("封闭环名称", "轴向间隙 X"),
    ("封闭环下限mm", 0.05),
    ("封闭环上限mm", 0.25),
    ("封闭环名义值mm", ""),
    ("目标公差mm", 0.20),
    ("分配方法", "it"),
    ("计算口径", "wc"),
    ("备注", "it=等精度法，equal=等公差法，proportional=等比法；wc=极值法，rss=概率法"),
]

HELP_TEXT = [
    "尺寸链与公差堆叠计算工具 —— Excel 输入模板",
    "",
    "填写顺序：",
    "1) 参数表：填写封闭环的验收上下限、目标公差与分配方法。",
    "2) 尺寸链表：每个组成环一行。类型填 增环 或 减环（增环使封闭环增大）。",
    "   上/下偏差按代数值填写，单位 mm；对称公差填 +0.05 / -0.05。",
    "   分布系数 k：正态 1.0（一般加工尺寸）、三角 1.22、均匀 1.73（形状误差）。",
    "   固定列填 是 表示该环公差不可分配（标准件、已有工艺能力）。",
    "3) GD&T 表：把几何公差标注换算成等效尺寸环，可留空。",
    "   作用列填 收紧（默认，几何误差使间隙变小）或 放宽。",
    "",
    "命令行：",
    "  python main.py calc   输入.xlsx        # 求解并输出计算书",
    "  python main.py alloc  输入.xlsx        # 反向公差分配",
    "  python main.py gdt    --feature 位置度 --value 0.2",
    "  python main.py template 输出.xlsx      # 重新生成此模板",
]


def _import_openpyxl():
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("需要 openpyxl：pip install openpyxl") from exc
    return dict(Workbook=Workbook, load_workbook=load_workbook, Alignment=Alignment,
                Border=Border, Font=Font, PatternFill=PatternFill, Side=Side,
                get_column_letter=get_column_letter, DataValidation=DataValidation)


def write_template(path: str, chain: Optional[Chain] = None,
                   gdt_rows: Optional[List[Dict[str, Any]]] = None) -> str:
    """生成（或用给定数据填充）一个 Excel 输入模板。"""
    ox = _import_openpyxl()
    wb = ox["Workbook"]()

    ws_help = wb.active
    ws_help.title = HELP_SHEET
    for i, line in enumerate(HELP_TEXT, start=1):
        ws_help.cell(row=i, column=1, value=line)
    ws_help.column_dimensions["A"].width = 100
    ws_help["A1"].font = ox["Font"](bold=True, size=13)

    ws_p = wb.create_sheet(PARAM_SHEET)
    ws_p.append(["参数", "取值"])
    src = _params_from_chain(chain) if chain else dict(PARAM_ROWS)
    for k, v in src.items():
        ws_p.append([k, v])
    _style_header(ws_p, 2, ox)
    ws_p.column_dimensions["A"].width = 20
    ws_p.column_dimensions["B"].width = 60

    ws_c = wb.create_sheet(CHAIN_SHEET)
    ws_c.append([h for h, _ in CHAIN_HEADERS])
    if chain:
        # GD&T 换算出来的环由 GD&T 表负责，写进尺寸链表会被读回时重复换算
        for i, r in enumerate([x for x in chain.rings if x.source != "GD&T"], start=1):
            ws_c.append([i, r.name, r.kind, r.nominal, r.es, r.ei, r.k, r.e,
                         "是" if r.fixed else "否", r.source, r.note])
    else:
        # 示例三环必须自洽：A0 = 20.00 - 1.00 - 18.90 = 0.10，落在 0.05~0.25 内
        ws_c.append([1, "A1 隔板厚度", "增环", 20.00, 0.05, 0.00, 1.0, 0.0,
                     "否", "尺寸", "示例：基孔制，下偏差 0"])
        ws_c.append([2, "A2 垫片厚度", "减环", 1.00, 0.02, -0.02, 1.0, 0.0,
                     "否", "尺寸", "示例：对称公差 ±0.02"])
        ws_c.append([3, "A3 轴套长度", "减环", 18.90, 0.00, -0.03, 1.0, 0.0,
                     "否", "尺寸", "示例：单向下偏差"])
    _style_header(ws_c, len(CHAIN_HEADERS), ox)
    for idx, (_, width) in enumerate(CHAIN_HEADERS, start=1):
        ws_c.column_dimensions[ox["get_column_letter"](idx)].width = width
    dv = ox["DataValidation"](type="list", formula1='"增环,减环"', allow_blank=False)
    ws_c.add_data_validation(dv)
    dv.add(f"C2:C400")

    ws_g = wb.create_sheet(GDT_SHEET)
    ws_g.append([h for h, _ in GDT_HEADERS])
    for i, row in enumerate(gdt_rows or [], start=1):
        ws_g.append([i, row.get("name", ""), row.get("feature", ""), row.get("value", ""),
                     row.get("kind", "减环"), row.get("mmc", "否"), row.get("size_tol", ""),
                     row.get("ref_size", ""),
                     row.get("bonus_mode", "mean"), row.get("effect", "收紧"),
                     row.get("note", "")])
    _style_header(ws_g, len(GDT_HEADERS), ox)
    for idx, (_, width) in enumerate(GDT_HEADERS, start=1):
        ws_g.column_dimensions[ox["get_column_letter"](idx)].width = width

    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    wb.save(path)
    return path


def _params_from_chain(chain: Chain) -> Dict[str, Any]:
    return {
        "项目名称": chain.title,
        "封闭环名称": chain.closed_name,
        "封闭环下限mm": '' if chain.limit_low is None else chain.limit_low,
        "封闭环上限mm": '' if chain.limit_high is None else chain.limit_high,
        "封闭环名义值mm": '' if chain.closed_nominal is None else chain.closed_nominal,
        "目标公差mm": '' if chain.target_tol is None else chain.target_tol,
        "分配方法": chain.alloc_method,
        "计算口径": chain.alloc_mode,
        "备注": chain.notes,
    }


def _style_header(ws, ncol: int, ox) -> None:
    fill = ox["PatternFill"]("solid", fgColor="DDEBF7")
    thin = ox["Side"](style="thin", color="B0B0B0")
    border = ox["Border"](left=thin, right=thin, top=thin, bottom=thin)
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = ox["Font"](bold=True)
        cell.fill = fill
        cell.alignment = ox["Alignment"](horizontal="center", vertical="center")
        cell.border = border
    ws.freeze_panes = "A2"


# ---------------------------------------------------------------- 读取
def read_chain(path: str) -> Chain:
    """按扩展名分派：.json 走 JSON 解析，.xlsx/.xlsm 走 Excel 解析。"""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".json":
        with open(path, "r", encoding="utf-8") as fh:
            return Chain.from_dict(json.load(fh))
    return read_excel(path)


def read_excel(path: str) -> Chain:
    ox = _import_openpyxl()
    wb = ox["load_workbook"](path, data_only=True)

    params: Dict[str, Any] = {}
    if PARAM_SHEET in wb.sheetnames:
        ws = wb[PARAM_SHEET]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[0] not in (None, ""):
                params[str(row[0]).strip()] = row[1] if len(row) > 1 else None

    if CHAIN_SHEET not in wb.sheetnames:
        raise ValueError(f"工作表缺失：{CHAIN_SHEET}")
    rings = _read_rings(wb[CHAIN_SHEET])

    gdt_rows: List[Dict[str, Any]] = []
    if GDT_SHEET in wb.sheetnames:
        gdt_rows = _read_gdt(wb[GDT_SHEET])
        rings.extend(_convert_gdt_rows(gdt_rows))

    if not rings:
        raise ValueError("尺寸链表为空：至少填写一个组成环")

    chain = Chain(
        title=_s(params.get("项目名称")) or os.path.basename(path),
        rings=rings,
        closed_name=_s(params.get("封闭环名称")) or "封闭环",
        limit_low=_num(params.get("封闭环下限mm")),
        limit_high=_num(params.get("封闭环上限mm")),
        closed_nominal=_num(params.get("封闭环名义值mm")),
        target_tol=_num(params.get("目标公差mm")),
        alloc_method=(_s(params.get("分配方法")) or "it").lower(),
        alloc_mode=(_s(params.get("计算口径")) or "wc").lower(),
        notes=_s(params.get("备注")) or "",
    )
    chain.gdt_rows = gdt_rows  # type: ignore[attr-defined]
    return chain


def _read_rings(ws) -> List[Ring]:
    rows = list(ws.iter_rows(values_only=True))
    header_idx = None
    for i, row in enumerate(rows[:15]):
        if row and any(_s(c) == "环名称" for c in row):
            header_idx = i
            break
    if header_idx is None:
        raise ValueError("在'尺寸链'表中找不到表头（需包含'环名称'列）")
    col = {_s(v): k for k, v in enumerate(rows[header_idx]) if _s(v)}

    def get(row, *names, default=None):
        for n in names:
            if n in col and col[n] < len(row):
                v = row[col[n]]
                if v not in (None, ""):
                    return v
        return default

    out: List[Ring] = []
    for row in rows[header_idx + 1:]:
        if not row or all(v in (None, "") for v in row):
            continue
        name = _s(get(row, "环名称"))
        if not name:
            continue
        kind = _s(get(row, "类型", default="增环")) or "增环"
        sign = -1 if ("减" in kind or kind.strip() in ("-", "-1")) else 1
        out.append(Ring(
            name=name,
            nominal=_num(get(row, "基本尺寸mm", "基本尺寸", default=0.0)) or 0.0,
            es=_num(get(row, "上偏差mm", "上偏差", default=0.0)) or 0.0,
            ei=_num(get(row, "下偏差mm", "下偏差", default=0.0)) or 0.0,
            sign=sign,
            k=_num(get(row, "分布系数k", "分布系数", default=1.0)) or 1.0,
            e=_num(get(row, "不对称系数e", "不对称系数", default=0.0)) or 0.0,
            fixed=_s(get(row, "固定", default="否")) in ("是", "Y", "y", "TRUE", "True"),
            source=_s(get(row, "来源", default="尺寸")) or "尺寸",
            note=_s(get(row, "备注", default="")) or "",
        ))
    return out


def _read_gdt(ws) -> List[Dict[str, Any]]:
    rows = list(ws.iter_rows(values_only=True))
    header_idx = None
    for i, row in enumerate(rows[:15]):
        if row and any(_s(c) == "几何特征" for c in row):
            header_idx = i
            break
    if header_idx is None:
        return []
    col = {_s(v): k for k, v in enumerate(rows[header_idx]) if _s(v)}

    def get(row, *names, default=None):
        for n in names:
            if n in col and col[n] < len(row):
                v = row[col[n]]
                if v not in (None, ""):
                    return v
        return default

    out = []
    for row in rows[header_idx + 1:]:
        if not row or all(v in (None, "") for v in row):
            continue
        feature = _s(get(row, "几何特征"))
        if not feature:
            continue
        kind = _s(get(row, "类型", default="减环")) or "减环"
        effect = _s(get(row, "作用", default="收紧")) or "收紧"
        out.append({
            "name": _s(get(row, "环名称")) or f"{feature}环",
            "feature": feature,
            "value": _num(get(row, "规格值mm", "规格值", default=0.0)) or 0.0,
            "sign": -1 if "减" in kind else 1,
            "mmc": _s(get(row, "是否Ⓜ", "是否M", default="否")) in ("是", "Y", "y"),
            "size_tol": _num(get(row, "尺寸公差mm", "尺寸公差", default=0.0)) or 0.0,
            "ref_size": _num(get(row, "要素尺寸mm", "要素尺寸", default=None)),
            "bonus_mode": (_s(get(row, "补偿方式", default="mean")) or "mean").lower(),
            "effect": "loosen" if ("放宽" in effect) else "tighten",
            "note": _s(get(row, "备注", default="")) or "",
        })
    return out


def _convert_gdt_rows(gdt_rows: List[Dict[str, Any]]) -> List[Ring]:
    from .gdttol import gdt_to_ring
    rings = []
    for row in gdt_rows:
        conv = gdt_to_ring(
            name=row["name"], feature=row["feature"], value=row["value"],
            sign=row["sign"], mmc=row["mmc"], size_tol=row["size_tol"],
            bonus_mode=row["bonus_mode"], effect=row["effect"],
            ref_size=row.get("ref_size"), note=row["note"])
        rings.append(conv.ring)
    return rings


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def _num(v) -> Optional[float]:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None
